"""Moliya: 🗓 10 kunlik jarimalar ro'yxati (amalga oshirildi / bekor qilish / orqaga)
va oy oxiridan N kun oldin keladigan moliya Excel hisoboti."""
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from conftest import run, sql


# ---------------- soxta Telegram obyektlari ----------------
class FakeBot:
    def __init__(self):
        self.sent, self.docs, self.edits = [], [], []

    async def send_message(self, chat_id, text, reply_markup=None, **kw):
        self.sent.append({"chat_id": chat_id, "text": text, "markup": reply_markup})

    async def send_document(self, chat_id, document, caption=None, **kw):
        self.docs.append({"chat_id": chat_id, "doc": document, "caption": caption})

    async def edit_message_text(self, text, chat_id=None, message_id=None,
                                reply_markup=None, **kw):
        self.edits.append({"text": text, "markup": reply_markup})

    async def edit_message_reply_markup(self, *a, **kw):
        return True


class FakeUser:
    def __init__(self, uid):
        self.id = uid
        self.full_name = "Test"


class FakeMessage:
    def __init__(self, tg_id, text=None):
        self.from_user = FakeUser(tg_id)
        self.chat = type("Chat", (), {"id": tg_id})()
        self.message_id = 777
        self.text = text
        self.answers, self.edits = [], []

    async def answer(self, text, reply_markup=None, **kw):
        self.answers.append({"text": text, "markup": reply_markup})

    async def edit_text(self, text, reply_markup=None, **kw):
        self.edits.append({"text": text, "markup": reply_markup})


class FakeCall:
    def __init__(self, tg_id, data):
        self.from_user = FakeUser(tg_id)
        self.data = data
        self.message = FakeMessage(tg_id)
        self.alerts = []

    async def answer(self, text=None, show_alert=False, **kw):
        self.alerts.append(text)


class FakeState:
    def __init__(self):
        self.state, self.data = None, {}

    async def set_state(self, st):
        self.state = st

    async def get_state(self):
        return self.state

    async def update_data(self, **kw):
        self.data.update(kw)

    async def get_data(self):
        return dict(self.data)

    async def clear(self):
        self.state, self.data = None, {}


def buttons(markup):
    if not markup:
        return []
    return [(b.text, b.callback_data) for row in markup.inline_keyboard for b in row]


FIN_TG, HR_TG = 820001, 820002


def _seed(env):
    sql(env, "INSERT INTO branches (id,name) VALUES (914,'Asaka filiali'), (915,'Soy filiali')")
    sql(env, "INSERT INTO users (id,tg_id,full_name,role) VALUES (?,?,?,?)",
        (8001, FIN_TG, "Moliya Xodimi", "accountant"))
    sql(env, "INSERT INTO users (id,tg_id,full_name,role) VALUES (?,?,?,?)",
        (8002, HR_TG, "HR Xodimi", "hr"))
    for uid, tg, name, br in ((8101, 830001, "Karimov Alisher Bahodirovich", 914),
                              (8102, 830002, "Ali Valiyev", 915),
                              (8103, 830003, "Toshmatova Nilufar", 914)):
        sql(env, "INSERT INTO users (id,tg_id,full_name,role,phone) VALUES (?,?,?,?,?)",
            (uid, tg, name, "pharmacist", "+998901234567"))
        sql(env, "INSERT INTO employee_profiles (user_id,role,position,branch_id) "
                 "VALUES (?,?,?,?)", (uid, "pharmacist", "Farmatsevt", br))


def _fine(env, fid, uid, amount, created, status="pending", branch=914):
    cancelled = 1 if status == "cancelled" else 0
    sql(env, "INSERT INTO fines (id,employee_user_id,branch_id,period,amount,reason,source,"
             "status,cancelled,created_by,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
        (fid, uid, branch, created[:7], amount, "Kechikish", "hr", status, cancelled,
         8002, created))


# ======================= YORDAMCHILAR =======================
def test_decade_helpers():
    from utils import (decade_start, decade_end, decade_prev, decade_next,
                       last_finished_decade, decade_tag, decade_from_tag)
    assert decade_start(date(2026, 10, 27)) == date(2026, 10, 21)
    assert decade_end(date(2026, 10, 21)) == date(2026, 10, 31)
    assert decade_end(date(2026, 2, 21)) == date(2026, 2, 28)
    assert decade_end(date(2026, 10, 11)) == date(2026, 10, 20)
    assert decade_prev(date(2026, 11, 1)) == date(2026, 10, 21)
    assert decade_next(date(2026, 10, 21)) == date(2026, 11, 1)
    assert last_finished_decade(date(2026, 10, 27)) == date(2026, 10, 11)
    assert last_finished_decade(date(2026, 11, 1)) == date(2026, 10, 21)
    assert decade_from_tag(decade_tag(date(2026, 10, 11))) == date(2026, 10, 11)
    assert decade_from_tag("261015") is None and decade_from_tag("xx") is None


def test_short_names():
    from utils import short_person_name, short_branch_name
    assert short_person_name("Karimov Alisher Bahodirovich") == "Karimov A."
    assert short_person_name("Ali Valiyev") == "Ali Valiyev"
    assert short_branch_name("Asaka filiali") == "Asaka f."
    assert short_branch_name("Yangi bozor filiali №2") == "Yangi bozor f. №2"
    assert short_branch_name(None) == "filialsiz"


# ======================= 10 KUNLIK RO'YXAT =======================
def test_digest_lists_people_with_branch_and_status(env):
    from handlers.fines import digest_view
    _seed(env)
    _fine(env, 1, 8101, "200 000 so'm", "2026-10-12 10:00:00")
    _fine(env, 2, 8101, "100 000 so'm", "2026-10-15 10:00:00", "executed")
    _fine(env, 3, 8102, "50 000 so'm", "2026-10-13 10:00:00", "executed", branch=915)
    _fine(env, 4, 8103, "70 000 so'm", "2026-10-14 10:00:00", "cancelled")
    _fine(env, 5, 8103, "90 000 so'm", "2026-10-05 10:00:00")  # oldingi davr

    text, markup = run(env, digest_view, date(2026, 10, 11), 0, date(2026, 10, 27))
    btns = buttons(markup)
    names = [t for t, c in btns if c.startswith("fdu:")]
    assert names == ["⏳ Karimov A. · Asaka f. · 300 000",
                     "✅ Ali Valiyev · Soy f. · 50 000",
                     "🚫 Toshmatova Nilufar · Asaka f. · bekor"]
    assert "amalga oshiring" in text
    assert "Kutilmoqda: <b>1</b> ta" in text
    assert "Oldingi davrlardan ijro kutilayotgan: <b>1</b> ta" in text
    cbs = [c for _, c in btns]
    assert "fdu:261011:8101:0" in cbs
    assert "fdl:261001:0" in cbs and "fdl:261021:0" in cbs  # oldingi / keyingi 10 kun


def test_card_execute_cancel_and_back(env):
    from handlers import fines as fh
    from states import FineDecisionForm
    _seed(env)
    _fine(env, 1, 8101, "200 000 so'm", "2026-10-12 10:00:00")
    _fine(env, 2, 8101, "100 000 so'm", "2026-10-15 10:00:00")
    bot = FakeBot()

    call = FakeCall(FIN_TG, "fdu:261011:8101:0")
    run(env, fh.fine_digest_card, call)
    card = call.message.edits[-1]
    assert "Karimov Alisher Bahodirovich" in card["text"]
    assert "Asaka filiali" in card["text"] and "300 000 so'm" in card["text"]
    cbs = [c for _, c in buttons(card["markup"])]
    assert cbs[-1] == "fdl:261011:0"  # ⬅️ Orqaga
    assert "fdx:261011:8101:1:0" in cbs and "fdc:261011:8101:2:0" in cbs
    assert "fdxa:261011:8101:0" in cbs

    # ✅ Amalga oshirildi
    call = FakeCall(FIN_TG, "fdx:261011:8101:1:0")
    run(env, fh.fine_digest_execute, call, bot)
    assert sql(env, "SELECT status FROM fines WHERE id=1")[0]["status"] == "executed"
    assert call.alerts[-1] == "✅ Jarima amalga oshirildi"
    assert any("ijro etildi" in m["text"] for m in bot.sent if m["chat_id"] == HR_TG)
    # ikkinchi marta bosilsa — xato xabari, holat o'zgarmaydi
    call = FakeCall(FIN_TG, "fdx:261011:8101:1:0")
    run(env, fh.fine_digest_execute, call, bot)
    assert "allaqachon" in call.alerts[-1]

    # 🚫 Bekor qilish — sabab so'raladi, keyin kartochka o'sha xabarda yangilanadi
    state = FakeState()
    call = FakeCall(FIN_TG, "fdc:261011:8101:2:0")
    run(env, fh.fine_digest_cancel, call, state)
    assert state.state == FineDecisionForm.cancel_reason
    msg = FakeMessage(FIN_TG, "Xato yozilgan")
    run(env, fh.fine_cancel_reason, msg, state, bot)
    row = sql(env, "SELECT status, cancel_reason FROM fines WHERE id=2")[0]
    assert row == {"status": "cancelled", "cancel_reason": "Xato yozilgan"}
    assert bot.edits and "Bekor qilindi" in bot.edits[-1]["text"]

    # ⬅️ Orqaga — ro'yxatga qaytadi
    call = FakeCall(FIN_TG, "fdl:261011:0")
    run(env, fh.fine_digest_list, call)
    assert "10 kunlik jarimalar" in call.message.edits[-1]["text"]


def test_execute_all_and_foreign_fine_rejected(env):
    from handlers import fines as fh
    _seed(env)
    _fine(env, 1, 8101, "200 000", "2026-10-12 10:00:00")
    _fine(env, 2, 8101, "100 000", "2026-10-15 10:00:00")
    _fine(env, 3, 8102, "50 000", "2026-10-13 10:00:00", branch=915)
    bot = FakeBot()
    # boshqa xodimning jarimasini shu kartochkadan ijro etib bo'lmaydi
    call = FakeCall(FIN_TG, "fdx:261011:8101:3:0")
    run(env, fh.fine_digest_execute, call, bot)
    assert sql(env, "SELECT status FROM fines WHERE id=3")[0]["status"] == "pending"
    call = FakeCall(FIN_TG, "fdxa:261011:8101:0")
    run(env, fh.fine_digest_execute_all, call, bot)
    assert call.alerts[-1] == "✅ 2 ta jarima amalga oshirildi"
    assert [r["status"] for r in sql(env, "SELECT status FROM fines ORDER BY id")] == [
        "executed", "executed", "pending"]
    # moliya bo'lmagan foydalanuvchi
    call = FakeCall(HR_TG, "fdxa:261011:8102:0")
    run(env, fh.fine_digest_execute_all, call, bot)
    assert sql(env, "SELECT status FROM fines WHERE id=3")[0]["status"] == "pending"


def test_card_with_many_long_fines_fits_telegram_limit(env):
    from handlers import fines as fh
    _seed(env)
    for i in range(1, 15):
        sql(env, "INSERT INTO fines (id,employee_user_id,branch_id,period,amount,reason,"
                 "source,status,created_by,created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
            (i, 8101, 914, "2026-10", "10 000", "<b>uzun</b> sabab " * 40, "hr",
             "pending", 8002, f"2026-10-12 10:{i:02d}:00"))
    text, markup = run(env, fh._card_view, date(2026, 10, 11), 8101, 0)
    assert len(text) <= fh.CARD_MAX_CHARS
    assert "yana" in text and "&lt;b&gt;uzun" in text  # HTML xavfsiz
    labels = [t for t, _ in buttons(markup)]
    assert "✅ Barchasini amalga oshirish (14 ta)" in labels


# ======================= AVTOMATIK YUBORISH =======================
def test_digest_sent_every_10_days_once(env):
    from services.reminders import _run_fine_digest
    _seed(env)
    _fine(env, 1, 8101, "200 000", "2026-10-12 10:00:00")
    bot = FakeBot()
    assert not run(env, _run_fine_digest, bot, datetime(2026, 10, 21, 8, 30))
    assert run(env, _run_fine_digest, bot, datetime(2026, 10, 21, 9, 5))
    assert [m["chat_id"] for m in bot.sent] == [FIN_TG]
    assert "11.10.2026 — 20.10.2026" in bot.sent[0]["text"]
    assert not run(env, _run_fine_digest, bot, datetime(2026, 10, 21, 12, 0))  # bir marta
    # oy oxiridagi davr (21–31) keyingi oyning 1-sanasida
    assert run(env, _run_fine_digest, bot, datetime(2026, 11, 1, 9, 0))
    assert "21.10.2026 — 31.10.2026" in bot.sent[-1]["text"]
    # juda kech (davrdan 3 kundan ko'p o'tgan) — eski ro'yxat yuborilmaydi
    assert not run(env, _run_fine_digest, bot, datetime(2026, 11, 8, 9, 0))


def test_finance_month_report_day_setting(env):
    from handlers.reports import finance_report_day
    from services.reminders import _run_finance_month_report
    assert finance_report_day(2026, 10, 10) == 21
    assert finance_report_day(2026, 10, 4) == 27
    assert finance_report_day(2026, 2, 0) == 28
    _seed(env)
    sql(env, "INSERT INTO settings (key,value) VALUES ('finance_report_days_before','10')")
    bot = FakeBot()
    assert not run(env, _run_finance_month_report, bot, datetime(2026, 10, 20, 12, 0))
    assert not run(env, _run_finance_month_report, bot, datetime(2026, 10, 21, 8, 0))
    assert run(env, _run_finance_month_report, bot, datetime(2026, 10, 21, 9, 30))
    assert [d["chat_id"] for d in bot.docs] == [FIN_TG]
    assert "Oy oxiriga 10 kun qoldi" in bot.docs[0]["caption"]
    assert not run(env, _run_finance_month_report, bot, datetime(2026, 10, 25, 9, 30))


def test_admin_sets_finance_report_days(env):
    from handlers import admin
    _seed(env)
    sql(env, "UPDATE users SET role='admin' WHERE id=8002")
    state = FakeState()
    msg = FakeMessage(HR_TG, "40")
    run(env, admin.finance_report_days_save, msg, state)
    assert "0 dan 27 gacha" in msg.answers[-1]["text"]
    msg = FakeMessage(HR_TG, "10")
    run(env, admin.finance_report_days_save, msg, state)
    assert sql(env, "SELECT value FROM settings WHERE key='finance_report_days_before'"
               )[0]["value"] == "10"
    assert "oy oxiridan <b>10</b> kun oldin" in msg.answers[-1]["text"]


# ======================= EXCEL =======================
def test_finance_excel_report_sheets_and_totals():
    from services.reports import DateRange, build_finance_report
    rng = DateRange(date(2026, 10, 1), date(2026, 10, 27), months=1)
    base = dict(employee_name="Karimov Alisher", branch_name="Asaka filiali",
                position="Farmatsevt", source="hr", created_by_name="HR Xodimi",
                reason="Kechikish")
    fines = [
        dict(base, id=1, employee_user_id=1, amount="200 000", status="pending",
             cancelled=0, created_at="2026-10-05 10:00:00"),
        dict(base, id=2, employee_user_id=1, amount="100 000", status="executed",
             cancelled=0, created_at="2026-10-15 10:00:00", executed_by_name="Moliya",
             executed_at="2026-10-16 09:00:00"),
        dict(base, id=3, employee_user_id=2, employee_name="Ali Valiyev",
             branch_name="Soy filiali", amount="50 000", status="cancelled", cancelled=1,
             cancel_reason="Xato", created_at="2026-10-22 10:00:00"),
    ]
    doc = build_finance_report(fines, rng)
    assert doc.filename.startswith("moliya_jarimalar_1oylik_")
    wb = load_workbook(BytesIO(doc.data))
    assert wb.sheetnames == ["Umumiy", "Xodimlar", "Jarimalar ro'yxati"]
    ws = wb["Umumiy"]
    assert ws["C5"].value == 300000          # oylikdan ushlanadi
    assert ws["G5"].value == 0.5             # amalga oshirish darajasi
    assert len(ws._charts) == 3
    # 10 kunliklar: [№, davr, soni, amalga oshirilgan, kutilmoqda, bekor, jami, xodimlar]
    decades = [list(r) for r in ws.iter_rows(min_row=12, max_row=15, values_only=True)]
    assert decades[0] == [1, "01.10 — 10.10.2026", 1, 0, 200000, 0, 200000, 1]
    assert decades[1] == [2, "11.10 — 20.10.2026", 1, 100000, 0, 0, 100000, 1]
    assert decades[2] == [3, "21.10 — 27.10.2026", 0, 0, 0, 50000, 0, 0]
    assert decades[3][1] == "JAMI" and decades[3][6] == 300000
    emp = wb["Xodimlar"]
    rows = {r[1]: r for r in emp.iter_rows(min_row=6, max_row=7, values_only=True)}
    assert rows["Karimov Alisher"][8] == 300000 and rows["Karimov Alisher"][9] == "Kutilmoqda"
    assert rows["Ali Valiyev"][8] == 0 and rows["Ali Valiyev"][9] == "Bekor qilingan"
    lst = wb["Jarimalar ro'yxati"]
    assert lst["C5"].value == "01.10 — 10.10" and lst["J7"].value == "Bekor qilingan"


def test_finance_excel_report_empty():
    from services.reports import DateRange, build_finance_report
    rng = DateRange(date(2026, 10, 1), date(2026, 10, 27), months=1)
    wb = load_workbook(BytesIO(build_finance_report([], rng).data))
    assert wb["Xodimlar"]["A6"].value == "Bu davrda jarima qo'llangan xodim yo'q"
