from io import BytesIO

import pytest
from openpyxl import load_workbook

from handlers.dayoff import _can_handle
from handlers.dayoff_plan import (
    canonical_weekday,
    dayoff_plan_mismatches,
    expected_day_status,
    plan_prompt_text,
    send_dayoff_report,
)
from services.export import build_dayoff_xlsx
from tests.conftest import run, sql


def test_weekday_normalization_and_expected_status():
    assert canonical_weekday("  Yakshanba ") == "Yakshanba"
    assert canonical_weekday("Воскресенье") == "Yakshanba"
    assert canonical_weekday("якшанба куни") == "Yakshanba"
    assert canonical_weekday("Суббота") == "Shanba"
    assert expected_day_status("Yakshanba", "Dushanba") == "work"
    assert expected_day_status("Понедельник", "Dushanba") == "off"
    assert expected_day_status("🚫 Dam olish kunim yo'q", "Dushanba") == "work"


def test_monday_weekend_employees_are_flagged_if_marked_off():
    plan = {
        "id": 7,
        "plan_date": "2026-10-05",
        "weekday": "Dushanba",
        "branch_name": "Paytug' filiali",
    }
    items = [
        {
            "id": 1,
            "full_name": "Зокирова Раънохон Мухторжоновна",
            "position": "Farmatsevt",
            "rest_day": "Yakshanba",
            "day_status": "off",
        },
        {
            "id": 2,
            "full_name": "Машрапов Самандар",
            "position": "Farmatsevt",
            "rest_day": "Shanba",
            "day_status": "off",
        },
    ]

    assert [item["id"] for item in dayoff_plan_mismatches(plan, items)] == [1, 2]
    text = plan_prompt_text(plan, items)
    assert text.count("JADVALDAN FARQ QILADI") == 2
    assert "Doimiy dam kuni: Yakshanba" in text
    assert "Doimiy dam kuni: Shanba" in text


def test_dayoff_excel_exposes_manual_override():
    report = build_dayoff_xlsx(
        [{
            "branch_name": "Paytug' filiali",
            "weekday": "Dushanba",
            "items": [{
                "full_name": "Машрапов Самандар",
                "position": "Farmatsevt",
                "rest_day": "Shanba",
                "day_status": "off",
                "scheduled_status": "work",
                "is_override": True,
            }],
        }],
        "05.10.2026",
    )
    ws = load_workbook(BytesIO(report.data), data_only=True).active
    rows = list(ws.iter_rows(values_only=True))

    assert rows[4] == (
        "№", "Ism-familiya", "Lavozim", "Doimiy dam kuni", "Holat", "Izoh"
    )
    assert rows[5][1:] == (
        "Машрапов Самандар", "Farmatsevt", "Shanba",
        "🛌 Dam oladi", "⚠️ Rahbar/HR o'zgartirgan",
    )


def test_plan_edit_can_be_reset_to_pending(env):
    sql(env, "INSERT INTO branches (id,name) VALUES (?,?)", (501, "Test filial"))
    sql(
        env,
        "INSERT INTO users (id,tg_id,full_name,role,branch_id) VALUES (?,?,?,?,?)",
        (601, 998000601, "Test xodim", "pharmacist", 501),
    )
    sql(
        env,
        "INSERT INTO employee_profiles (user_id,role,position,branch_id,rest_day) "
        "VALUES (?,?,?,?,?)",
        (601, "pharmacist", "Farmatsevt", 501, "Yakshanba"),
    )

    from database import queries as q

    plan_id = run(env, q.create_dayoff_plan, 501, "2026-10-05", "Dushanba")
    run(env, q.add_dayoff_plan_item, plan_id, 601, "Test xodim", "Farmatsevt", "work")
    run(env, q.set_dayoff_plan_status, plan_id, "confirmed", 601)
    item = run(env, q.list_dayoff_plan_items, plan_id)[0]
    assert item["rest_day"] == "Yakshanba"

    run(env, q.toggle_dayoff_plan_item, item["id"])
    plan = run(env, q.get_dayoff_plan, plan_id)
    assert plan["status"] == "pending"
    assert plan["confirmed_by"] is None
    assert plan["confirmed_at"] is None


@pytest.mark.asyncio
async def test_manager_cannot_handle_another_branch_request():
    manager = {"id": 10, "role": "manager", "branch_id": 100}
    assert await _can_handle(manager, {"branch_id": 100}) is True
    assert await _can_handle(manager, {"branch_id": 200}) is False
    assert await _can_handle({"id": 11, "role": "hr"}, {"branch_id": 200}) is True


@pytest.mark.asyncio
async def test_official_report_excludes_pending_branches(monkeypatch):
    from handlers import dayoff_plan as module

    async def plans(_date, status=None):
        if status == "confirmed":
            return [{
                "id": 1, "branch_name": "Tasdiqlangan filial",
                "weekday": "Dushanba", "status": "confirmed",
            }]
        return [{
            "id": 2, "branch_name": "Kutilayotgan filial",
            "weekday": "Dushanba", "status": "pending",
        }]

    async def items(plan_id):
        assert plan_id == 1
        return [{
            "id": 1, "full_name": "Test xodim", "position": "Farmatsevt",
            "rest_day": "Dushanba", "day_status": "off",
        }]

    monkeypatch.setattr(module.q, "list_dayoff_plans_for_date", plans)
    monkeypatch.setattr(module.q, "list_dayoff_plan_items", items)

    class Bot:
        def __init__(self):
            self.documents = []

        async def send_document(self, chat_id, document, caption=None):
            self.documents.append((chat_id, document, caption))

    bot = Bot()
    assert await send_dayoff_report(bot, [123], "2026-10-05") == 1
    _, document, caption = bot.documents[0]
    ws = load_workbook(BytesIO(document.data), data_only=True).active
    all_text = " ".join(
        str(value)
        for row in ws.iter_rows(values_only=True)
        for value in row if value is not None
    )
    assert "Tasdiqlangan filial" in all_text
    assert "Kutilayotgan filial" not in all_text
    assert "Kutilayotgan filial" in caption


def test_confirmed_plan_text_marks_overrides():
    from handlers.dayoff_plan import confirmed_plan_text
    plan = {"plan_date": "2026-10-09", "weekday": "Juma",
            "branch_name": "Paytug' filiali", "confirmed_at": "2026-10-08 17:06:00"}
    items = [
        {"id": 1, "full_name": "Ulmasov Ahmadullo", "rest_day": "Juma", "day_status": "off"},
        {"id": 2, "full_name": "Зокирова Раънохон", "rest_day": "Yakshanba", "day_status": "off"},
        {"id": 3, "full_name": "Шамсутдинов", "rest_day": "Жума", "day_status": "work"},
    ]
    text = confirmed_plan_text(plan, items, "Солийев Султонбек")
    assert "Tasdiqladi: Солийев Султонбек" in text
    assert "Ulmasov Ahmadullo — Juma" in text
    assert "doimiy dam kuni: Yakshanba ⚠️ jadvaldan farqli" in text
    # Kirillcha dam kuni kanonik nomga keltiriladi va «keladi» farqi ko'rinadi
    assert "Шамсутдинов — dam kuni bo'lsa ham ishga keladi" in text
    assert "farqli o'zgarish: <b>2</b>" in text


def test_hr_status_list_shows_confirmed_and_pending(env):
    from database import queries as q
    from handlers import dayoff_plan as module
    for bid, name in ((701, "Asaka filiali"), (702, "Qo'rg'ontepa filiali"),
                      (703, "Yangi filial")):
        sql(env, "INSERT INTO branches (id,name) VALUES (?,?)", (bid, name))
    # Rahbar: users.branch_id bo'sh, filial faqat profilda
    sql(env, "INSERT INTO users (id,tg_id,full_name,role) VALUES (?,?,?,?)",
        (801, 998000801, "Rahbar Profil", "manager"))
    sql(env, "INSERT INTO employee_profiles (user_id,role,position,branch_id) "
             "VALUES (?,?,?,?)", (801, "manager", "Filial rahbari", 702))
    p1 = run(env, q.create_dayoff_plan, 701, "2026-10-09", "Juma")
    p2 = run(env, q.create_dayoff_plan, 702, "2026-10-09", "Juma")
    run(env, q.set_dayoff_plan_status, p1, "confirmed", None)

    text, markup = run(env, module._status_list_view, "2026-10-09")
    labels = [b.text for row in markup.inline_keyboard for b in row]
    datas = [b.callback_data for row in markup.inline_keyboard for b in row]
    assert "✅ Tasdiqlagan: <b>1</b>" in text
    assert "❌ Tasdiqlamagan: <b>1</b>" in text
    assert labels[0] == "❌ Qo'rg'ontepa filiali"  # tasdiqlamaganlar tepada
    assert any(lbl.startswith("✅ Asaka filiali") for lbl in labels)
    assert any(lbl.startswith("⚪ Yangi filial") for lbl in labels)
    assert f"dolbr:{p2}" in datas

    managers = run(env, module.branch_manager_tg_ids, 702)
    assert managers == [998000801]
